import streamlit as st
import pandas as pd
import os
import hashlib
from datetime import datetime, timedelta, timezone
from PIL import Image
import io

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# FUSO HORÁRIO DE BRASÍLIA (UTC-3)
FUSO_BR = timezone(timedelta(hours=-3))

def agora_brasil():
    return datetime.now(FUSO_BR).strftime('%d/%m/%Y %H:%M')

# CARREGAMENTO DA LOGO
ARQUIVO_LOGO = "logo.png"
logo_img = None

if os.path.exists(ARQUIVO_LOGO):
    try:
        logo_img = Image.open(ARQUIVO_LOGO)
    except Exception:
        logo_img = None

# CONFIGURAÇÃO DA PÁGINA
st.set_page_config(
    page_title="Copa Ambiental - Manutenção", 
    page_icon=logo_img if logo_img else "🚛", 
    layout="wide"
)

# CSS MINIMALISTA E LIMPO
estilo_limpo = """
    <style>
    #MainMenu {visibility: hidden;}
    footer {visibility: hidden;}
    header {visibility: hidden;}
    [data-testid="stHeader"] {display: none;}
    
    div.stButton > button {
        width: 100%;
        height: 48px;
        font-size: 16px !important;
        font-weight: 600 !important;
        border-radius: 8px !important;
        border: 1px solid #e0e0e0 !important;
        transition: all 0.2s ease;
    }
    div.stButton > button:hover {
        border-color: #008000 !important;
        color: #008000 !important;
    }
    </style>
"""
st.markdown(estilo_limpo, unsafe_allow_html=True)

# ARQUIVOS DE BANCO DE DADOS
ARQUIVO_CSV = 'chamados_manutencao.csv'
ARQUIVO_USUARIOS = 'usuarios.csv'
ARQUIVO_PAPEL_TIMBRADO = '8. Papel Timbrado.docx'

VEICULOS = [
    "Caminhão Compactador", "Caminhão Poliguindaste", "Caminhão Roll-On",
    "Caminhão Pipa", "Caminhão Basculante", "Carregadeira",
    "Retroescavadeira", "Trator de Esteira", "Motoniveladora",
    "Pick-up Operacional", "Van de Equipe", "Veículo Leve / Apoio"
]

def hash_senha(senha):
    return hashlib.sha256(str.encode(senha)).hexdigest()

def carregar_usuarios():
    if not os.path.exists(ARQUIVO_USUARIOS):
        df = pd.DataFrame(columns=['usuario', 'senha', 'nome', 'nivel'])
        df.to_csv(ARQUIVO_USUARIOS, index=False)
        return df
    
    df = pd.read_csv(ARQUIVO_USUARIOS)
    df['nivel'] = df['nivel'].astype(float)
    return df

def salvar_usuario(novo_user, nova_senha, nome, nivel, nivel_criador):
    df = carregar_usuarios()
    nome_limpo = nome.strip()
    if " " not in nome_limpo:
        return False, "Por favor, digite o Nome e o Sobrenome completo!"

    if float(nivel) >= 4.0:
        return False, "Não é permitido cadastrar novos usuários de nível 4.0!"

    if novo_user in df['usuario'].values:
        return False, "Usuário já existe!"
    
    senha_h = hash_senha(nova_senha)
    if senha_h in df['senha'].values:
        return False, "Esta senha já está em uso por outro usuário."
    
    novo_df = pd.DataFrame([{
        'usuario': novo_user,
        'senha': senha_h,
        'nome': nome_limpo,
        'nivel': float(nivel)
    }])
    df = pd.concat([df, novo_df], ignore_index=True)
    df.to_csv(ARQUIVO_USUARIOS, index=False)
    return True, "Usuário cadastrado com sucesso!"

def atualizar_nivel_usuario(user_alvo, novo_nivel, nivel_editor, user_logado):
    if user_alvo == user_logado:
        return False, "Você não pode alterar o seu próprio nível de acesso!"

    df = carregar_usuarios()
    if user_alvo in df['usuario'].values:
        nivel_alvo = float(df.loc[df['usuario'] == user_alvo, 'nivel'].values[0])
        if float(novo_nivel) >= 4.0 or nivel_alvo >= 4.0:
            return False, "Não é permitido alterar níveis de Administrador Global (4.0)!"
        if float(nivel_editor) <= nivel_alvo:
            return False, "Você não tem permissão para alterar o nível deste usuário!"
            
        df.loc[df['usuario'] == user_alvo, 'nivel'] = float(novo_nivel)
        df.to_csv(ARQUIVO_USUARIOS, index=False)
        return True, f"Nível atualizado com sucesso!"
    return False, "Usuário não encontrado!"

def atualizar_nome_usuario(user_alvo, novo_nome, nivel_editor):
    if float(nivel_editor) < 3.5:
        return False, "Sem permissão!"
    nome_limpo = novo_nome.strip()
    if " " not in nome_limpo:
        return False, "Digite nome e sobrenome!"

    df = carregar_usuarios()
    if user_alvo in df['usuario'].values:
        df.loc[df['usuario'] == user_alvo, 'nome'] = nome_limpo
        df.to_csv(ARQUIVO_USUARIOS, index=False)
        return True, "Nome atualizado!"
    return False, "Usuário não encontrado!"

def redefinir_senha_usuario(user_alvo, nova_senha, nivel_editor, user_logado):
    df = carregar_usuarios()
    if user_alvo in df['usuario'].values:
        senha_h = hash_senha(nova_senha)
        df.loc[df['usuario'] == user_alvo, 'senha'] = senha_h
        df.to_csv(ARQUIVO_USUARIOS, index=False)
        return True, "Senha alterada!"
    return False, "Usuário não encontrado!"

def excluir_usuario(user_alvo, user_logado, nivel_editor):
    if user_alvo == user_logado:
        return False, "Não pode excluir sua própria conta!"
    df = carregar_usuarios()
    if user_alvo in df['usuario'].values:
        df = df[df['usuario'] != user_alvo]
        df.to_csv(ARQUIVO_USUARIOS, index=False)
        return True, "Usuário excluído!"
    return False, "Usuário não encontrado!"

def limpar_chamados_expirados(df):
    if df.empty or 'Data' not in df.columns:
        return df
    agora = datetime.now(FUSO_BR)
    indices_para_remover = []
    for idx, row in df.iterrows():
        if str(row.get('Aprovado_Coordenador', 'Não')).strip() == 'Não':
            try:
                data_chamado = datetime.strptime(str(row['Data']), '%d/%m/%Y %H:%M').replace(tzinfo=FUSO_BR)
                if agora - data_chamado > timedelta(days=7):
                    indices_para_remover.append(idx)
            except Exception:
                pass
    if indices_para_remover:
        df = df.drop(index=indices_para_remover).reset_index(drop=True)
        salvar_dados(df)
    return df

def carregar_dados():
    colunas_obrigatorias = [
        'ID_OS', 'Data', 'Motorista', 'Veiculo', 'Placa', 
        'Descricao_Problema', 'Status', 'Prioridade', 'Aprovado_Coordenador', 
        'Data_Aprovacao', 'Mecanico_Responsavel', 'Data_Liberacao', 'Arquivado'
    ]
    if not os.path.exists(ARQUIVO_CSV):
        df = pd.DataFrame(columns=colunas_obrigatorias)
        df.to_csv(ARQUIVO_CSV, index=False)
        return df
    
    df = pd.read_csv(ARQUIVO_CSV)
    for col in colunas_obrigatorias:
        if col not in df.columns:
            df[col] = ''
            
    df = df.fillna({
        'Motorista': 'Não Identificado',
        'Descricao_Problema': 'Sem descrição',
        'Mecanico_Responsavel': 'Não Atribuído',
        'Prioridade': 'Média',
        'Arquivado': 'Não',
        'Status': 'Aguardando Aprovação',
        'Data_Aprovacao': '',
        'Data_Liberacao': ''
    })
    df = limpar_chamados_expirados(df)
    salvar_dados(df)
    return df

def salvar_dados(df):
    df.to_csv(ARQUIVO_CSV, index=False)

def gerar_relatorio_word(df_rel, subtitulo_filtro=""):
    # Carrega o modelo de papel timbrado existente para preservar cabeçalhos, rodapés e formatação padrão
    if os.path.exists(ARQUIVO_PAPEL_TIMBRADO):
        try:
            doc = Document(ARQUIVO_PAPEL_TIMBRADO)
        except Exception:
            doc = Document()
    else:
        doc = Document()
    
    # Adiciona o título no corpo do documento carregado
    p_titulo = doc.add_paragraph()
    p_titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_tit = p_titulo.add_run("RELATÓRIO DE MANUTENÇÃO DO VEÍCULO")
    run_tit.bold = True
    run_tit.font.size = Pt(14)
    run_tit.font.color.rgb = RGBColor(0, 100, 0)
    
    if subtitulo_filtro:
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_sub = p_sub.add_run(subtitulo_filtro)
        run_sub.font.size = Pt(11)
        run_sub.font.italic = True
    
    p_data_geracao = doc.add_paragraph()
    p_data_geracao.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_data_geracao.add_run(f"Emitido em: {agora_brasil()}")
    
    doc.add_paragraph()

    def formatar_data_hora(val):
        if not val or str(val).lower() == 'nan' or str(val).strip() == '':
            return 'Não registrada'
        val_str = str(val).strip()
        try:
            dt = datetime.strptime(val_str, '%d/%m/%Y %H:%M')
            return dt.strftime('%d/%m/%Y %H:%M')
        except ValueError:
            try:
                dt = datetime.fromisoformat(val_str)
                return dt.strftime('%d/%m/%Y %H:%M')
            except Exception:
                return val_str

    estilos_disponiveis = [s.name for s in doc.styles]
    sub_style = 'List Bullet 2' if 'List Bullet 2' in estilos_disponiveis else 'List Bullet'

    # Adiciona cada chamado utilizando listas com marcadores limpos e sem duplicidade no corpo do timbrado
    for _, row in df_rel.iterrows():
        p_os_titulo = doc.add_paragraph(style='List Bullet')
        run_os_num = p_os_titulo.add_run(f"Ordem de Serviço: {str(row.get('ID_OS', ''))}")
        run_os_num.bold = True
        
        doc.add_paragraph(f"Veículo / Equipamento: {str(row.get('Veiculo', ''))}", style=sub_style)
        doc.add_paragraph(f"Identificação / Placa: {str(row.get('Placa', ''))}", style=sub_style)
        doc.add_paragraph(f"Data do Registro: {formatar_data_hora(row.get('Data', ''))}", style=sub_style)
        doc.add_paragraph(f"Data de Aprovação: {formatar_data_hora(row.get('Data_Aprovacao', ''))}", style=sub_style)
        doc.add_paragraph(f"Prioridade: {str(row.get('Prioridade', ''))}", style=sub_style)
        doc.add_paragraph(f"Mecânico Responsável: {str(row.get('Mecanico_Responsavel', ''))}", style=sub_style)
        doc.add_paragraph(f"Data de Liberação: {formatar_data_hora(row.get('Data_Liberacao', ''))}", style=sub_style)
        doc.add_paragraph(f"Descrição do Problema: {str(row.get('Descricao_Problema', ''))}", style=sub_style)
        
        doc.add_paragraph() # Espaçador entre os chamados

    f = io.BytesIO()
    doc.save(f)
    f.seek(0)
    return f

# CONTROLE DE SESSÃO
if 'logged_in' not in st.session_state:
    st.session_state['logged_in'] = False
    st.session_state['user_info'] = None

if 'aba_ativa' not in st.session_state:
    st.session_state['aba_ativa'] = 'Menu'

# TELA DE LOGIN
if not st.session_state['logged_in']:
    col_l1, col_l2, col_l3 = st.columns([1, 2, 1])
    with col_l2:
        if logo_img:
            st.image(logo_img, width=280)
        st.subheader("Acesso ao Sistema")
        
        with st.form("form_login"):
            user_input = st.text_input("Usuário").strip().lower()
            password_input = st.text_input("Senha", type="password")
            btn_login = st.form_submit_button("Entrar", use_container_width=True)
            
            if btn_login:
                df_users = carregar_usuarios()
                senha_enc = hash_senha(password_input)
                user_match = df_users[(df_users['usuario'] == user_input) & (df_users['senha'] == senha_enc)]
                
                if not user_match.empty:
                    st.session_state['logged_in'] = True
                    st.session_state['user_info'] = user_match.iloc[0].to_dict()
                    st.session_state['aba_ativa'] = 'Menu'
                    st.rerun()
                else:
                    st.error("Usuário ou senha incorretos.")

else:
    user_data = st.session_state['user_info']
    nivel_user = float(user_data['nivel'])
    usuario_atual = str(user_data['usuario'])
    lbl_nivel = str(int(nivel_user)) if nivel_user.is_integer() else str(nivel_user)

    if logo_img:
        st.sidebar.image(logo_img, use_container_width=True)
    st.sidebar.write(f"👤 **{user_data['nome']}**")
    st.sidebar.caption(f"Nível de Acesso: {lbl_nivel}")
    st.sidebar.markdown("---")
    
    if st.sidebar.button("🏠 Menu Principal", key="sb_home", use_container_width=True):
        st.session_state['aba_ativa'] = 'Menu'
        st.rerun()
        
    if st.sidebar.button("🚪 Sair / Logout", key="sb_logout", use_container_width=True):
        st.session_state['logged_in'] = False
        st.session_state['user_info'] = None
        st.session_state['aba_ativa'] = 'Menu'
        st.rerun()

    df_os = carregar_dados()

    if st.session_state['aba_ativa'] == 'Menu':
        st.title("Menu Principal")
        st.caption(f"Bem-vindo, {user_data['nome']}")
        
        opcoes = [
            ("📝 Abrir Chamado", "Abrir Chamado"),
            ("🔍 Consultar Chamados", "Consultar Chamados")
        ]
        
        if nivel_user == 2.0 or nivel_user >= 4.0:
            opcoes.append(("🛠️ Painel da Oficina", "Oficina"))
            
        if nivel_user >= 3.0:
            opcoes.append(("🎯 Triagem e Prioridade", "Triagem"))
            opcoes.append(("👤 Gestão de Usuários", "Usuarios"))

        opcoes.append(("🚪 Sair / Logout", "Logout"))

        col_m1, col_m2 = st.columns(2)
        for idx, (label, chave) in enumerate(opcoes):
            coluna = col_m1 if idx % 2 == 0 else col_m2
            with coluna:
                if st.button(label, key=f"btn_menu_{chave}", use_container_width=True):
                    if chave == "Logout":
                        st.session_state['logged_in'] = False
                        st.session_state['user_info'] = None
                        st.session_state['aba_ativa'] = 'Menu'
                        st.rerun()
                    else:
                        st.session_state['aba_ativa'] = chave
                        st.rerun()

    else:
        col_voltar, col_titulo = st.columns([1, 4])
        with col_voltar:
            if st.button("← Voltar"):
                st.session_state['aba_ativa'] = 'Menu'
                st.rerun()

        # PÁGINA 1: ABRIR CHAMADO
        if st.session_state['aba_ativa'] == "Abrir Chamado":
            st.header("📝 Nova Ordem de Serviço")
            col_veic, col_chk = st.columns([3, 1])
            with col_veic:
                veiculo_sel = st.selectbox("Veículo / Equipamento", VEICULOS)
            with col_chk:
                st.write("")
                st.write("")
                outro_marcado = st.checkbox("Outros")

            outros_veiculo = ""
            if outro_marcado:
                outros_veiculo = st.text_input("Especifique o veículo/equipamento")

            with st.form("form_chamado", clear_on_submit=True):
                placa = st.text_input("Placa ou Identificação").upper()
                descricao = st.text_area("Descrição do Problema / Defeito")
                btn_submeter = st.form_submit_button("Enviar Chamado", use_container_width=True)

                if btn_submeter:
                    veiculo_final = outros_veiculo if outro_marcado else veiculo_sel
                    if outro_marcado and not outros_veiculo.strip():
                        st.warning("Por favor, especifique o veículo.")
                    elif placa and descricao:
                        novo_id = f"OS-{len(df_os) + 1001}"
                        nova_os = {
                            'ID_OS': novo_id,
                            'Data': agora_brasil(),
                            'Motorista': user_data['nome'],
                            'Veiculo': veiculo_final,
                            'Placa': placa,
                            'Descricao_Problema': descricao,
                            'Status': 'Aguardando Aprovação',
                            'Prioridade': 'Pendente',
                            'Aprovado_Coordenador': 'Não',
                            'Data_Aprovacao': '',
                            'Mecanico_Responsavel': 'Não Atribuído',
                            'Data_Liberacao': '',
                            'Arquivado': 'Não'
                        }
                        df_os = pd.concat([df_os, pd.DataFrame([nova_os])], ignore_index=True)
                        salvar_dados(df_os)
                        st.success(f"Chamado {novo_id} enviado com sucesso!")
                    else:
                        st.warning("Preencha a placa e a descrição.")

        # PÁGINA 2: CONSULTAR CHAMADOS
        elif st.session_state['aba_ativa'] == "Consultar Chamados":
            st.header("🔍 Consultar Ordens de Serviço")
            
            c1, c2 = st.columns([3, 1])
            with c1:
                busca_placa = st.text_input("Buscar por Placa ou Identificação").upper()
            with c2:
                ver_arquivados = st.selectbox("Exibir", ["Ativos", "Arquivados"])
            
            df_exibicao = df_os.copy()
            status_arq = "Sim" if ver_arquivados == "Arquivados" else "Não"
            df_exibicao = df_exibicao[df_exibicao['Arquivado'] == status_arq]
            
            if nivel_user == 1.0:
                colunas_nivel_1 = ['ID_OS', 'Data', 'Veiculo', 'Placa', 'Descricao_Problema', 'Status']
                df_exibicao = df_exibicao[colunas_nivel_1]
                df_exibicao['Status'] = df_exibicao['Status'].replace({
                    'Aguardando Aprovação': 'Em Aberto',
                    'Aguardando Manutenção': 'Em Aberto'
                })
                df_exibicao.columns = ['Nº OS', 'Data', 'Veículo', 'Placa', 'Descrição', 'Status']
            else:
                colunas_gestao = ['ID_OS', 'Data', 'Motorista', 'Veiculo', 'Placa', 'Descricao_Problema', 'Status', 'Prioridade', 'Mecanico_Responsavel']
                df_exibicao = df_exibicao[colunas_gestao]
                df_exibicao.columns = ['Nº OS', 'Data', 'Solicitante', 'Veículo', 'Placa', 'Descrição', 'Status', 'Prioridade', 'Mecânico']

            if busca_placa:
                df_exibicao = df_exibicao[df_exibicao['Placa'].astype(str).str.contains(busca_placa, na=False)]

            st.dataframe(df_exibicao, use_container_width=True)

            # Seção de Relatório Individual por Veículo/Equipamento para Níveis 3.0+
            if nivel_user >= 3.0:
                st.markdown("")
                with st.container(border=True):
                    st.subheader("📥 Baixar Relatório Individual (Papel Timbrado)")
                    
                    df_rel_base = df_os.copy()
                    if status_arq == "Sim":
                        df_rel_base = df_rel_base[df_rel_base['Arquivado'] == 'Sim']
                    else:
                        df_rel_base = df_rel_base[df_rel_base['Arquivado'] != 'Sim']

                    if busca_placa:
                        df_rel_base = df_rel_base[df_rel_base['Placa'].astype(str).str.contains(busca_placa, na=False)]

                    veiculos_disponiveis = df_rel_base['Veiculo'].dropna().unique().tolist()
                    
                    if not veiculos_disponiveis:
                        st.info("Nenhum veículo disponível para exportação com os filtros atuais.")
                    else:
                        col_sel_v, col_btn_v = st.columns([2, 1])
                        with col_sel_v:
                            veiculo_escolhido = st.selectbox("Selecione o Veículo/Equipamento", veiculos_disponiveis)
                        
                        with col_btn_v:
                            st.write("")
                            df_veiculo_especifico = df_rel_base[df_rel_base['Veiculo'] == veiculo_escolhido]
                            arquivo_docx = gerar_relatorio_word(df_veiculo_especifico, subtitulo_filtro=f"Veículo / Equipamento: {veiculo_escolhido}")

                            st.download_button(
                                label=f"Baixar ({veiculo_escolhido})",
                                data=arquivo_docx,
                                file_name=f"relatorio_manutencao_{veiculo_escolhido.replace(' ', '_').lower()}_{datetime.now(FUSO_BR).strftime('%Y%m%d_%H%M')}.docx",
                                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                                use_container_width=True
                            )

            if nivel_user >= 4.0:
                st.markdown("")
                with st.container(border=True):
                    st.subheader("⚙️ Gestão de OS (Administrador Global)")
                    lista_os = df_os['ID_OS'].tolist()
                    if lista_os:
                        os_selecionada = st.selectbox("Selecione a OS para Ação", lista_os)
                        dados_os_sel = df_os[df_os['ID_OS'] == os_selecionada].iloc[0]
                        idx_os = df_os[df_os['ID_OS'] == os_selecionada].index[0]

                        col_arq, col_del = st.columns(2)
                        with col_arq:
                            status_atual_arq = dados_os_sel.get('Arquivado', 'Não')
                            lbl_btn = "Desarquivar" if status_atual_arq == "Sim" else "Arquivar"
                            if st.button(lbl_btn, use_container_width=True):
                                df_os.at[idx_os, 'Arquivado'] = "Não" if status_atual_arq == "Sim" else "Sim"
                                salvar_dados(df_os)
                                st.success("Status atualizado!")
                                st.rerun()

                        with col_del:
                            if st.button("🗑️ Excluir Definitivamente", type="primary", use_container_width=True):
                                df_os = df_os.drop(index=idx_os).reset_index(drop=True)
                                salvar_dados(df_os)
                                st.success("OS excluída!")
                                st.rerun()

        # PÁGINA 3: TRIAGEM
        elif st.session_state['aba_ativa'] == "Triagem":
            if nivel_user < 3.0:
                st.error("Acesso não autorizado!")
            else:
                st.header("🎯 Triagem & Prioridades")
                pendentes = df_os[(df_os['Aprovado_Coordenador'] == 'Não') & (df_os['Arquivado'] != 'Sim')]

                if pendentes.empty:
                    st.info("Nenhum chamado pendente de aprovação.")
                else:
                    for idx, row in pendentes.iterrows():
                        with st.expander(f"{row['ID_OS']} - {row['Veiculo']} ({row['Placa']})"):
                            st.write(f"**Solicitante:** {row['Motorista']} | **Data:** {row['Data']}")
                            st.write(f"**Problema:** {row['Descricao_Problema']}")
                            
                            prioridade = st.selectbox(f"Prioridade", ["Alta", "Média", "Baixa"], key=f"prio_{idx}")
                            if st.button(f"Aprovar e Enviar para Oficina", key=f"btn_aprov_{idx}", use_container_width=True):
                                df_os.at[idx, 'Aprovado_Coordenador'] = 'Sim'
                                df_os.at[idx, 'Prioridade'] = prioridade
                                df_os.at[idx, 'Status'] = 'Aguardando Manutenção'
                                df_os.at[idx, 'Data_Aprovacao'] = agora_brasil()
                                salvar_dados(df_os)
                                st.success(f"{row['ID_OS']} aprovada!")
                                st.rerun()

        # PÁGINA 4: OFICINA
        elif st.session_state['aba_ativa'] == "Oficina":
            if nivel_user != 2.0 and nivel_user < 4.0:
                st.error("Acesso não autorizado!")
            else:
                st.header("🛠️ Painel da Oficina")
                aprovados = df_os[(df_os['Aprovado_Coordenador'] == 'Sim') & (df_os['Arquivado'] != 'Sim')]
                
                aba_pend, aba_and, aba_conc, aba_busca = st.tabs(["⏳ Em Aberto", "🔄 Em Andamento", "✅ Concluídos Recentes", "🔍 Filtrar Histórico"])
                
                def renderizar_cards(df_sub, aba_nome):
                    if df_sub.empty:
                        st.info(f"Nenhum chamado em '{aba_nome}'.")
                    else:
                        for idx, row in df_sub.iterrows():
                            with st.expander(f"[{row['Prioridade']}] {row['ID_OS']} - {row['Veiculo']} ({row['Placa']})"):
                                st.write(f"**Solicitante:** {row['Motorista']} | **Data do Registro:** {row['Data']}")
                                st.write(f"**Problema:** {row['Descricao_Problema']}")
                                
                                novo_status = st.selectbox("Status", ["Aguardando Manutenção", "Em Andamento", "Concluído"], index=["Aguardando Manutenção", "Em Andamento", "Concluído"].index(row['Status']) if row['Status'] in ["Aguardando Manutenção", "Em Andamento", "Concluído"] else 0, key=f"st_{idx}")
                                
                                mec_atual = str(row['Mecanico_Responsavel']).strip()
                                mec_bloqueado = mec_atual and mec_atual != 'Não Atribuído' and mec_atual != ''
                                
                                if mec_bloqueado:
                                    st.text_input("Mecânico Responsável", value=mec_atual, disabled=True, key=f"mec_dis_{idx}")
                                    mecanico = mec_atual
                                    st.caption("🔒 Mecânico bloqueado após registro inicial.")
                                else:
                                    mecanico = st.text_input("Mecânico Responsável", value="", key=f"mec_{idx}")
                                
                                chave_confirma = f"confirma_mec_{idx}"
                                if chave_confirma not in st.session_state:
                                    st.session_state[chave_confirma] = False

                                if st.button(f"Salvar Alterações", key=f"btn_m_{idx}", use_container_width=True):
                                    if not mec_bloqueado:
                                        if not mecanico or mecanico.strip() == '' or mecanico == 'Não Atribuído':
                                            st.warning("Informe o nome do mecânico responsável.")
                                        else:
                                            st.session_state[chave_confirma] = True
                                            st.session_state[f"temp_mec_{idx}"] = mecanico.strip()
                                            st.session_state[f"temp_status_{idx}"] = novo_status
                                            st.rerun()
                                    else:
                                        df_os.at[idx, 'Status'] = novo_status
                                        if novo_status == 'Concluído' and not str(row['Data_Liberacao']).strip():
                                            df_os.at[idx, 'Data_Liberacao'] = agora_brasil()
                                        elif novo_status != 'Concluído':
                                            df_os.at[idx, 'Data_Liberacao'] = ''
                                        salvar_dados(df_os)
                                        st.success("Atualizado com sucesso!")
                                        st.rerun()

                                if st.session_state.get(chave_confirma, False):
                                    nome_temp = st.session_state[f"temp_mec_{idx}"]
                                    st.warning(f"⚠️ **Atribuir {nome_temp} a função?** (Esta ação bloqueará a alteração futura do mecânico).")
                                    col_conf1, col_conf2 = st.columns(2)
                                    with col_conf1:
                                        if st.button("✅ Sim, Confirmar", key=f"sim_{idx}", use_container_width=True):
                                            df_os.at[idx, 'Status'] = st.session_state[f"temp_status_{idx}"]
                                            df_os.at[idx, 'Mecanico_Responsavel'] = nome_temp
                                            
                                            st_fin = st.session_state[f"temp_status_{idx}"]
                                            if st_fin == 'Concluído' and not str(row['Data_Liberacao']).strip():
                                                df_os.at[idx, 'Data_Liberacao'] = agora_brasil()
                                            
                                            salvar_dados(df_os)
                                            st.session_state[chave_confirma] = False
                                            st.success("Mecânico atribuído e bloqueado com sucesso!")
                                            st.rerun()
                                    with col_conf2:
                                        if st.button("❌ Cancelar", key=f"nao_{idx}", use_container_width=True):
                                            st.session_state[chave_confirma] = False
                                            st.rerun()

                with aba_pend:
                    renderizar_cards(aprovados[aprovados['Status'] == 'Aguardando Manutenção'], "Em Aberto")
                with aba_and:
                    renderizar_cards(aprovados[aprovados['Status'] == 'Em Andamento'], "Em Andamento")
                with aba_conc:
                    def eh_recente(data_str):
                        try:
                            dt = datetime.strptime(str(data_str).strip(), '%d/%m/%Y %H:%M')
                            agora = datetime.now(FUSO_BR).replace(tzinfo=None)
                            return (agora - dt).days <= 2
                        except Exception:
                            return False
                    
                    concluidos_todos = aprovados[aprovados['Status'] == 'Concluído']
                    concluidos_recentes = concluidos_todos[concluidos_todos['Data_Liberacao'].apply(eh_recente)]
                    renderizar_cards(concluidos_recentes, "Concluídos Recentes")

                with aba_busca:
                    st.subheader("🔍 Filtro Avançado de Chamados")
                    col_f1, col_f2, col_f3, col_f4 = st.columns(4)
                    with col_f1:
                        f_status = st.selectbox("Status", ["Todos", "Aguardando Aprovação", "Aguardando Manutenção", "Em Andamento", "Concluído"], key="f_status")
                    with col_f2:
                        f_mecanico = st.text_input("Mecânico Responsável", key="f_mec")
                    with col_f3:
                        f_data = st.text_input("Data (DD/MM/AAAA)", key="f_data")
                    with col_f4:
                        f_placa = st.text_input("Placa / ID", key="f_placa")

                    df_filtrado = df_os[df_os['Arquivado'] != 'Sim'].copy()
                    if f_status != "Todos":
                        df_filtrado = df_filtrado[df_filtrado['Status'] == f_status]
                    if f_mecanico.strip():
                        df_filtrado = df_filtrado[df_filtrado['Mecanico_Responsavel'].astype(str).str.contains(f_mecanico, case=False, na=False)]
                    if f_data.strip():
                        df_filtrado = df_filtrado[
                            df_filtrado['Data'].astype(str).str.contains(f_data, case=False, na=False) | 
                            df_filtrado['Data_Aprovacao'].astype(str).str.contains(f_data, case=False, na=False) | 
                            df_filtrado['Data_Liberacao'].astype(str).str.contains(f_data, case=False, na=False)
                        ]
                    if f_placa.strip():
                        df_filtrado = df_filtrado[
                            df_filtrado['Placa'].astype(str).str.contains(f_placa, case=False, na=False) | 
                            df_filtrado['ID_OS'].astype(str).str.contains(f_placa, case=False, na=False)
                        ]

                    st.markdown("")
                    st.dataframe(
                 
