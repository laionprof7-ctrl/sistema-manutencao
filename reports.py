import io
import os
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from config import ARQUIVO_PAPEL_TIMBRADO
from database import agora_brasil


def gerar_relatorio_word(df_rel, subtitulo_filtro=""):
    if os.path.exists(ARQUIVO_PAPEL_TIMBRADO):
        try:
            doc = Document(ARQUIVO_PAPEL_TIMBRADO)
        except Exception:
            doc = Document()
    else:
        doc = Document()

    p_titulo = doc.add_paragraph()
    p_titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_titulo.add_run("RELATÓRIO DE MANUTENÇÃO DO VEÍCULO")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0, 100, 0)

    if subtitulo_filtro:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(subtitulo_filtro)
        r.font.size = Pt(11)
        r.font.italic = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run(f"Emitido em: {agora_brasil()}")
    doc.add_paragraph()

    def fmt(valor):
        txt = str(valor or "").strip()
        if not txt or txt.lower() == "nan":
            return "Não registrada"
        try:
            return datetime.strptime(txt, "%d/%m/%Y %H:%M").strftime("%d/%m/%Y %H:%M")
        except ValueError:
            return txt

    estilos = [s.name for s in doc.styles]
    sub_style = "List Bullet 2" if "List Bullet 2" in estilos else "List Bullet"

    for _, row in df_rel.iterrows():
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(f"Ordem de Serviço: {row.get('ID_OS', '')}")
        r.bold = True
        campos = [
            ("Veículo / Equipamento", row.get("Veiculo", "")),
            ("Identificação / Placa", row.get("Placa", "")),
            ("Data do Registro", fmt(row.get("Data", ""))),
            ("Data de Aprovação", fmt(row.get("Data_Aprovacao", ""))),
            ("Prioridade", row.get("Prioridade", "")),
            ("Mecânico Responsável", row.get("Mecanico_Responsavel", "")),
            ("Data de Liberação", fmt(row.get("Data_Liberacao", ""))),
            ("Descrição do Problema", row.get("Descricao_Problema", "")),
        ]
        for nome, valor in campos:
            doc.add_paragraph(f"{nome}: {valor}", style=sub_style)
        doc.add_paragraph()

    f = io.BytesIO()
    doc.save(f)
    f.seek(0)
    return f
